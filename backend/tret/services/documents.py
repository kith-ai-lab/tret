"""Document text extraction (PDF/DOCX/XLSX/PPTX/CSV/MD/TXT) and the shared
ingestion step — bounded extraction plus building the `Document` row — that
both `api/documents.py`'s upload endpoint and the connections import endpoint
(`POST /api/projects/{project_id}/documents/import`) call, so the two paths
can never diverge on how a document actually gets ingested once its bytes
are on disk. Only how the bytes arrived differs between them.
"""
from __future__ import annotations

import asyncio
import csv
import io
import logging
import re
import uuid
from pathlib import Path

from tret.db.models import Document

log = logging.getLogger("tret.documents")

# The one size cap for every path that ends up in `ingest_document` — a
# manual upload (`api/documents.py::upload_document`, which streams the body
# and enforces this mid-stream) and a provider import (`api/documents.py::
# _import_one`, via `services/connections.py::IMPORT_MAX_BYTES`, which holds
# the whole downloaded file in memory before it ever reaches here). 25MB is
# sized to the import path's memory cost, not the upload path's — streaming
# could afford a much larger cap, but importing cannot, and both paths must
# share one number rather than silently diverge on what "too large" means.
MAX_DOCUMENT_BYTES = 25 * 1024 * 1024

# Extraction bounds. A 25MB CSV of one-character rows, or a PDF crafted to
# expand, must not turn into an unbounded string in a JSONB column or an
# unbounded stretch of CPU inside a request.
MAX_EXTRACTED_CHARS = 2_000_000
EXTRACTION_TIMEOUT_SECONDS = 30.0

# .xlsx and .pptx are zip containers that openpyxl/python-pptx inflate while
# parsing. A small compressed file can declare an enormous uncompressed size
# per member (a zip bomb), so the total is summed straight from the zip's
# central directory — no inflation, no parsing — before either library ever
# touches the bytes.
#
# .xlsx keeps the full 200MB allowance: `_extract_xlsx` reads `read_only`,
# streaming row by row rather than building an in-memory object model, so a
# large-but-legitimate workbook's memory cost stays bounded regardless of
# how much the zip expands to. `_extract_pptx` has no such streaming path —
# `python-pptx`'s `Presentation()` parses the whole package eagerly into
# memory before extraction ever gets to bound anything — so a pptx gets a
# lower allowance instead of trusting the same 200MB ceiling to be safe for
# a parser that cannot stream.
MAX_ZIP_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_PPTX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024


def _refuse_zip_bombs(filename: str, data: bytes, max_uncompressed_bytes: int = MAX_ZIP_UNCOMPRESSED_BYTES) -> None:
    import zipfile

    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        total = sum(info.file_size for info in zf.infolist())
    if total > max_uncompressed_bytes:
        raise ValueError(
            f"{filename}: archive expands to {total} bytes, over the "
            f"{max_uncompressed_bytes} byte limit"
        )


def extract_text(filename: str, data: bytes) -> tuple[str, dict]:
    """Returns (text, meta). Raises ValueError for unsupported types."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        # A `## Page N` marker ahead of every page — the same heading-marker
        # convention `_extract_xlsx`/`_extract_pptx` already use for sheets and
        # slides — is what gives `services/retrieval.py`'s chunker a page
        # locator to attach to prose that otherwise has no heading structure
        # of its own (pypdf returns flowing text, not a document tree).
        parts = [f"## Page {i}\n{page.extract_text() or ''}" for i, page in enumerate(reader.pages, start=1)]
        return "\n\n".join(parts), {"pages": len(reader.pages)}
    if lower.endswith(".docx"):
        return _extract_docx(data)
    if lower.endswith(".xlsx"):
        return _extract_xlsx(filename, data)
    if lower.endswith(".pptx"):
        return _extract_pptx(filename, data)
    if lower.endswith((".csv", ".tsv")):
        delim = "\t" if lower.endswith(".tsv") else ","
        text = data.decode("utf-8", errors="replace")
        rows = list(csv.reader(io.StringIO(text), delimiter=delim))
        rendered = "\n".join(" | ".join(row) for row in rows)
        return rendered, {"rows": len(rows), "columns": rows[0] if rows else []}
    if lower.endswith((".md", ".txt", ".json", ".yaml", ".yml")):
        return data.decode("utf-8", errors="replace"), {}
    raise ValueError(
        f"Unsupported file type: {filename} (supported: pdf, docx, xlsx, pptx, csv, tsv, md, txt)"
    )


_HEADING_STYLE_RE = re.compile(r"^Heading (\d+)$")


def _docx_heading_level(style_name: str) -> int:
    """0 for a body-text style, 1-6 for a heading style ("Title" counts as a
    top-level heading same as "Heading 1"; "Heading 7"+ clamps to 6, matching
    the deepest markdown heading `services/retrieval.py`'s chunker recognises).
    """
    if style_name == "Title":
        return 1
    m = _HEADING_STYLE_RE.match(style_name)
    return min(6, int(m.group(1))) if m else 0


def _extract_docx(data: bytes) -> tuple[str, dict]:
    """Paragraphs and tables, in document order, with headings and tables
    marked inline (`# heading text` / `## Table N`) rather than dumped as flat
    text and every table appended after every paragraph the way python-docx's
    `.paragraphs`/`.tables` accessors would — document order and heading
    structure are what let `services/retrieval.py`'s chunker build heading
    paths and keep tables whole instead of treating the whole file as one
    undifferentiated block of prose.

    Blocks (one per paragraph, and one per table — heading plus every row) are
    joined with a blank line between them, not a single newline: a docx has no
    blank-line convention of its own (each paragraph is already its own XML
    element), but `services/retrieval.py`'s prose chunker finds paragraph
    breaks by splitting on blank lines, same as it does for a `.md` upload. A
    single-newline join would hand it one run-on paragraph per section and
    default straight to the hard-split path instead of packing paragraphs to
    the ~250-400 token target. A table's own rows stay single-newline-joined
    *within* the table's one block, exactly like every other extractor's
    table/sheet rows — only the blank line between blocks is new.
    """
    import docx
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(data))
    blocks: list[str] = []
    table_count = 0
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, document)
            text = paragraph.text
            if not text.strip():
                continue
            level = _docx_heading_level((paragraph.style.name if paragraph.style else "") or "")
            blocks.append(f"{'#' * level} {text}" if level else text)
        elif isinstance(child, CT_Tbl):
            table_count += 1
            table = Table(child, document)
            rows = [" | ".join(cell.text for cell in row.cells) for row in table.rows]
            blocks.append("\n".join([f"## Table {table_count}", *rows]))
    return "\n\n".join(blocks), {"paragraphs": len(document.paragraphs), "tables": table_count}


def _extract_xlsx(filename: str, data: bytes) -> tuple[str, dict]:
    """Row-major text of every worksheet, via `read_only` iteration so a 25MB
    workbook with millions of cells is streamed rather than built into an
    in-memory object model. Formulas are never evaluated: `data_only=True`
    reads each cell's last cached value (what Excel wrote when it last saved
    the file), not a live recomputation.
    """
    import openpyxl

    _refuse_zip_bombs(filename, data)
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet_count = len(workbook.sheetnames)
        parts: list[str] = []
        chars = 0
        total_rows = 0
        truncated = False
        for sheet in workbook.worksheets:
            if truncated:
                break
            heading = f"## Sheet: {sheet.title}"
            parts.append(heading)
            chars += len(heading) + 1
            for row in sheet.iter_rows(values_only=True):
                if row is None or all(value is None for value in row):
                    continue
                line = "\t".join("" if value is None else str(value) for value in row)
                parts.append(line)
                chars += len(line) + 1
                total_rows += 1
                if chars >= MAX_EXTRACTED_CHARS:
                    truncated = True
                    break
    finally:
        workbook.close()
    meta: dict = {"sheets": sheet_count, "rows": total_rows}
    if truncated:
        meta["truncated"] = True
    return "\n".join(parts), meta


def _extract_pptx(filename: str, data: bytes) -> tuple[str, dict]:
    """Per-slide text: every shape's text frame (paragraph per line), table
    cells tab-separated, and speaker notes under a `Notes:` heading.
    """
    from pptx import Presentation

    _refuse_zip_bombs(filename, data, MAX_PPTX_UNCOMPRESSED_BYTES)
    presentation = Presentation(io.BytesIO(data))
    parts: list[str] = []
    chars = 0
    truncated = False
    for index, slide in enumerate(presentation.slides, start=1):
        if truncated:
            break
        heading = f"## Slide {index}"
        parts.append(heading)
        chars += len(heading) + 1
        for shape in slide.shapes:
            if truncated:
                break
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    if not paragraph.text:
                        continue
                    parts.append(paragraph.text)
                    chars += len(paragraph.text) + 1
                    if chars >= MAX_EXTRACTED_CHARS:
                        truncated = True
                        break
            if truncated:
                break
            if shape.has_table:
                for row in shape.table.rows:
                    line = "\t".join(cell.text for cell in row.cells)
                    parts.append(line)
                    chars += len(line) + 1
                    if chars >= MAX_EXTRACTED_CHARS:
                        truncated = True
                        break
        if truncated:
            break
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            notes_text = notes_frame.text if notes_frame else ""
            if notes_text.strip():
                parts.append("Notes:")
                parts.append(notes_text)
                chars += len(notes_text) + 8
                if chars >= MAX_EXTRACTED_CHARS:
                    truncated = True
    meta: dict = {"slides": len(presentation.slides)}
    if truncated:
        meta["truncated"] = True
    return "\n".join(parts), meta


async def extract_bounded(filename: str, data: bytes) -> tuple[str, dict, str]:
    """(text, meta, status). Bounded in wall time and in characters returned.

    Extraction is third-party parsing of third-party bytes: it runs off the event
    loop so one slow document cannot stall every other request, it is abandoned
    after EXTRACTION_TIMEOUT_SECONDS, its output is truncated, and *any* parser
    failure is recorded on the row rather than raised — a corrupt PDF is a
    document with no text, not a 500 on upload.

    "Abandoned" is precise: Python cannot cancel a running thread, so a pathological
    parser finishes in the background while the request moves on. What is bounded is
    the request, and the memory that thread holds (one document, at most the size cap
    its caller enforced before getting here).
    """
    try:
        text, meta = await asyncio.wait_for(
            asyncio.to_thread(extract_text, filename, data), EXTRACTION_TIMEOUT_SECONDS
        )
    except (asyncio.TimeoutError, TimeoutError):
        return "", {"error": f"text extraction timed out after {EXTRACTION_TIMEOUT_SECONDS:.0f}s"}, "failed"
    except ValueError as e:  # unsupported type: the message is for the user
        return "", {"error": str(e)}, "failed"
    except Exception as e:  # a parser blowing up on hostile bytes
        log.warning("text extraction failed for %r: %s", filename, e)
        return "", {"error": f"text extraction failed: {type(e).__name__}"}, "failed"
    if len(text) > MAX_EXTRACTED_CHARS:
        meta = {**meta, "truncated": True, "extracted_chars": MAX_EXTRACTED_CHARS}
        text = text[:MAX_EXTRACTED_CHARS]
    return text, meta, "done"


async def ingest_document(
    *,
    project_id: uuid.UUID,
    filename: str,
    content_type: str,
    storage_path: Path,
    sha256: str,
    byte_size: int,
    data: bytes,
    uploaded_by: uuid.UUID | None,
    source_kind: str = "upload",
    meta_extra: dict | None = None,
) -> Document:
    """Extract `data`'s text (bounded) and build the `Document` row for it.

    `data` is already written to `storage_path` by the caller — an upload
    streams it there while consuming an untrusted multipart body, an import
    downloads it whole from a provider — so this is purely "bytes in, a row
    ready to add to the session out". Not committed and not added to the
    session: the caller owns the transaction (the upload endpoint commits
    once; the import endpoint commits per item so one bad file in a batch
    cannot roll back the others).
    """
    text, meta, status = await extract_bounded(filename, data)
    if meta_extra:
        meta = {**meta, **meta_extra}
    return Document(
        project_id=project_id,
        filename=filename,
        content_type=content_type,
        byte_size=byte_size,
        storage_path=str(storage_path),
        extracted_text=text or None,  # None = nothing was extracted
        extraction_status=status,
        meta=meta,
        sha256=sha256,
        source_kind=source_kind,
        uploaded_by=uploaded_by,
    )
